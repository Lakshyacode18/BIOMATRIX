"""Export model assumptions for qualified-clinician review."""
import csv
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill

from knowledge import DISEASES, FEATURES, PRIORS


REVIEW_NOTES = {
    "__calibration__": (
        "Review probability calibration on an independent set of published, de-identified cases; "
        "the current model percentages may be overconfident."
    ),
    "__prior__": (
        "Review against the intended season, region, and clinic type; this is a synthetic starting assumption."
    ),
    "no_response_paracetamol": (
        "Review whether antipyretic response meaningfully separates these diseases; fever reduction may be temporary."
    ),
    "relative_bradycardia": (
        "Review the Typhoid value against clinical evidence; this sign may occur in only a minority of patients."
    ),
}


def generate_review_sheet(path="review_sheet.csv", workbook_path="review_sheet.xlsx"):
    rows = [[
        "disease",
        "feature",
        "kind",
        "model_value",
        "doctor_value",
        "source",
        "comment",
    ]]
    rows.append([
        "All diseases",
        "Probability calibration",
        "model_review",
        "Uncalibrated",
        "",
        "",
        REVIEW_NOTES["__calibration__"],
    ])
    for disease in DISEASES:
        rows.append([
            disease,
            "Assumed disease prior",
            "prior",
            PRIORS[disease],
            "",
            "",
            REVIEW_NOTES["__prior__"],
        ])
        for feature, details in FEATURES.items():
            rows.append([
                disease,
                details["label"],
                details["kind"],
                details["p"][disease],
                "",
                "",
                REVIEW_NOTES.get(feature, ""),
            ])

    with open(path, "w", newline="", encoding="utf-8") as file:
        csv.writer(file).writerows(rows)

    workbook = Workbook()
    review = workbook.active
    review.title = "Review"
    review.append([
        "Disease", "Key", "Kind", "Our estimate", "Doctor value",
        "Decision", "Source", "Comment",
    ])
    for disease, feature, kind, model_value, _, _, comment in rows[1:]:
        if kind == "prior":
            key = "Assumed disease prior"
        elif kind == "model_review":
            key = feature
        else:
            key = next(key for key, details in FEATURES.items() if details["label"] == feature)
        review.append([disease, key, kind, model_value, None, None, None, comment])
    review_decisions = DataValidation(type="list", formula1='"Change,Agree,Unsure"', allow_blank=True)
    review.add_data_validation(review_decisions)
    review_decisions.add(f"F2:F{review.max_row}")

    priors = workbook.create_sheet("Priors")
    priors.append(["Clinic type", "Indian outpatient fever workup"])
    priors.append([])
    priors.append(["Disease", "Our prior", "Doctor prior", "Decision", "Source", "Comment"])
    for disease in DISEASES:
        priors.append([disease, PRIORS[disease], None, None, None, REVIEW_NOTES["__prior__"]])
    priors.append(["Total", "=SUM(B4:B11)", "=SUM(C4:C11)", None, None, None])
    prior_decisions = DataValidation(type="list", formula1='"Change,Agree,Unsure"', allow_blank=True)
    priors.add_data_validation(prior_decisions)
    prior_decisions.add(f"D4:D{priors.max_row - 1}")

    header_fill = PatternFill("solid", fgColor="DDE8EF")
    for sheet in (review, priors):
        sheet.freeze_panes = "A2" if sheet is review else "A4"
        for cell in sheet[1 if sheet is review else 3]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 14), 58)
            sheet.column_dimensions[letter].width = width

    for row in review.iter_rows(min_row=2):
        if row[2].value in {"feature", "prior"}:
            row[3].number_format = "0.0%"
            row[4].number_format = "0.0%"
    for row in priors.iter_rows(min_row=4, max_col=3):
        row[1].number_format = "0.0%"
        row[2].number_format = "0.0%"
    workbook.save(workbook_path)
    print(f"Wrote {path} and {workbook_path}.")


if __name__ == "__main__":
    generate_review_sheet()